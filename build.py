#!/usr/bin/env python3
"""Сборка презентации VITOMEK из корпоративного шаблона.

    python3 build.py content.json -o result.pptx            # проверить и собрать
    python3 build.py content.json -o result.pptx --render   # + картинки слайдов и проверка переполнения
    python3 build.py --check content.json                   # только проверить content.json
    python3 build.py --catalog                              # показать макеты и лимиты

Нужен только python-pptx (pip install python-pptx). Для --render — LibreOffice (soffice) и poppler
(pdftoppm, pdftotext); без них сборка работает, проверка по картинкам пропускается.
"""
import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit('Нужен python-pptx: pip install python-pptx')

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / 'template' / 'vitomek-template.pptx'
CATALOG = json.loads((HERE / 'layouts.json').read_text(encoding='utf-8'))
LAYOUTS = CATALOG['layouts']


# ------------------------------------------------------------------ проверка содержания

def _items(value):
    """list-поле -> [(text, level)]"""
    if isinstance(value, str):
        value = [v for v in value.split('\n') if v.strip()]
    out = []
    for v in value or []:
        if isinstance(v, dict):
            out.append((str(v.get('text', '')), int(v.get('level', 0))))
        else:
            out.append((str(v), 0))
    return out


def validate(spec):
    """-> (errors, warnings). errors — собирать нельзя, warnings — соберётся, но стоит поправить."""
    errors, warns = [], []
    slides = spec.get('slides') if isinstance(spec, dict) else None
    if not slides:
        return ['В content.json нет списка "slides"'], []
    for n, s in enumerate(slides, 1):
        lay = s.get('layout')
        if lay not in LAYOUTS:
            errors.append(f'Слайд {n}: неизвестный макет «{lay}». Есть: {", ".join(LAYOUTS)}')
            continue
        cfg = LAYOUTS[lay]
        known = set(cfg['fields']) | {'layout', 'notes'} | ({cfg['items']['key']} if 'items' in cfg else set())
        for k in s:
            if k not in known:
                errors.append(f'Слайд {n} («{lay}»): у макета нет поля "{k}". Поля: {", ".join(sorted(known - {"layout"}))}')
        for key, f in cfg['fields'].items():
            v = s.get(key)
            if v in (None, '', []):
                if f.get('required') and 'default' not in f:
                    errors.append(f'Слайд {n} («{lay}»): не заполнено обязательное поле "{key}" ({f["label"]})')
                continue
            if f['type'] == 'text':
                if len(str(v)) > f['max']:
                    warns.append(f'Слайд {n} («{lay}») {key}: {len(str(v))} зн. при лимите {f["max"]} — сократи')
            elif f['type'] == 'list':
                items = _items(v)
                if len(items) > f['max_items']:
                    errors.append(f'Слайд {n} («{lay}») {key}: {len(items)} пунктов при максимуме {f["max_items"]} — '
                                  f'сократи или разбей на два слайда')
                for t, _ in items:
                    if len(t) > f['max']:
                        warns.append(f'Слайд {n} («{lay}») {key}: пункт «{t[:40]}…» {len(t)} зн. при лимите {f["max"]}')
                total = sum(len(t) for t, _ in items)
                if f.get('max_total') and total > f['max_total']:
                    warns.append(f'Слайд {n} («{lay}») {key}: всего {total} зн. при лимите {f["max_total"]}')
            elif f['type'] == 'image':
                if not Path(str(v)).exists():
                    warns.append(f'Слайд {n} («{lay}») image: файл «{v}» не найден — поле останется для ручной вставки')
            elif f['type'] == 'table':
                hdr, rows = v.get('header') or [], v.get('rows') or []
                ncol = max([len(hdr)] + [len(r) for r in rows]) if (hdr or rows) else 0
                if ncol > f['max_cols']:
                    errors.append(f'Слайд {n} таблица: {ncol} столбцов при максимуме {f["max_cols"]}')
                if len(rows) + (1 if hdr else 0) > f['max_rows']:
                    errors.append(f'Слайд {n} таблица: {len(rows) + (1 if hdr else 0)} строк при максимуме {f["max_rows"]} — '
                                  f'раздели на два слайда')
                for r in [hdr] + rows:
                    for c in r:
                        if len(str(c)) > f['max']:
                            warns.append(f'Слайд {n} таблица: ячейка «{str(c)[:30]}…» {len(str(c))} зн. при лимите {f["max"]}')
        if 'items' in cfg:
            it = cfg['items']
            vals = s.get(it['key']) or []
            if not (it['min'] <= len(vals) <= it['max']):
                hint = (' Для двух пунктов возьми «Две колонки», для 4+ — «Заголовок и текст».'
                        if lay == 'Три карточки' else '')
                errors.append(f'Слайд {n} («{lay}»): {it["key"]} — нужно от {it["min"]} до {it["max"]}, '
                              f'а передано {len(vals)}.{hint}')
            for i, v in enumerate(vals[:it['max']], 1):
                for fk, f in it['fields'].items():
                    t = str((v or {}).get(fk, ''))
                    if not t and f.get('required'):
                        errors.append(f'Слайд {n} {it["key"]}[{i}]: не заполнено "{fk}" ({f["label"]})')
                    elif len(t) > f['max']:
                        warns.append(f'Слайд {n} {it["key"]}[{i}].{fk}: {len(t)} зн. при лимите {f["max"]} — сократи')
    if slides and slides[0].get('layout') != 'Обложка':
        warns.append('Первый слайд обычно «Обложка»')
    if slides and slides[-1].get('layout') != 'Спасибо за внимание':
        warns.append('Последний слайд обычно «Спасибо за внимание»')
    run = 1
    for a, b in zip(slides, slides[1:]):
        run = run + 1 if a.get('layout') == b.get('layout') else 1
        if run == 4:
            warns.append(f'Больше трёх слайдов подряд на макете «{b.get("layout")}» — чередуй подачу')
    return errors, warns


# ------------------------------------------------------------------ сборка

def _layout(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    sys.exit(f'В шаблоне нет макета «{name}»')


def _set_text(ph, value, as_list=False):
    tf = ph.text_frame
    items = _items(value) if as_list else [(str(value), 0)]
    # оставляем один абзац, остальные удаляем — оформление наследуется от макета
    txBody = tf._txBody
    for p in txBody.findall(qn('a:p'))[1:]:
        txBody.remove(p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p0._p.remove(r)
    for i, (text, level) in enumerate(items):
        p = p0 if i == 0 else tf.add_paragraph()
        p.level = max(0, min(level, 4))
        lines = text.split('\n')
        for li, line in enumerate(lines):
            if li:
                p.add_line_break()
            p.add_run().text = line


def _hide_placeholder(ph):
    el = ph._element
    el.find('.//' + qn('p:cNvPr')).set('hidden', '1')
    sppr = el.find(qn('p:spPr'))
    if sppr is not None:
        for tag in ('a:noFill', 'a:solidFill', 'a:ln'):
            for x in sppr.findall(qn(tag)):
                sppr.remove(x)
        sppr.append(sppr.makeelement(qn('a:noFill'), {}))
        ln = sppr.makeelement(qn('a:ln'), {})
        ln.append(ln.makeelement(qn('a:noFill'), {}))
        sppr.append(ln)
    if ph.has_text_frame:
        _set_text(ph, ' ')


def _add_slide_number(slide, layout):
    for ph in layout.placeholders:
        if ph.placeholder_format.type is not None and 'SLIDE_NUMBER' in str(ph.placeholder_format.type):
            el = copy.deepcopy(ph._element)
            slide.shapes._spTree.append(el)
            return


def _table_from_example(prs_originals, layout_name):
    for s in prs_originals:
        if s.slide_layout.name == layout_name:
            for sh in s.shapes:
                if sh.has_table:
                    return sh._element
    return None


def _fill_table(frame_el, header, rows):
    """Подогнать таблицу-образец под нужное число строк/столбцов и заполнить, сохраняя оформление ячеек."""
    tbl = frame_el.find('.//' + qn('a:tbl'))
    grid = tbl.find(qn('a:tblGrid'))
    trs = tbl.findall(qn('a:tr'))
    data = ([header] if header else []) + rows
    ncol = max(len(r) for r in data)
    # столбцы
    cols = grid.findall(qn('a:gridCol'))
    total_w = sum(int(c.get('w')) for c in cols)
    while len(cols) < ncol:
        nc = copy.deepcopy(cols[-1]); grid.append(nc); cols.append(nc)
    while len(cols) > ncol:
        grid.remove(cols.pop())
    for c in cols:
        c.set('w', str(total_w // ncol))
    for tr in trs:
        tcs = tr.findall(qn('a:tc'))
        while len(tcs) < ncol:
            nt = copy.deepcopy(tcs[-1]); tr.append(nt); tcs.append(nt)
        while len(tcs) > ncol:
            tr.remove(tcs.pop())
    # строки: образец шапки и двух чередующихся строк тела
    head_tpl = trs[0]
    body_tpl = trs[1:3] if len(trs) >= 3 else trs[1:2] or trs[:1]
    for tr in trs:
        tbl.remove(tr)
    new_rows = []
    if header:
        new_rows.append((copy.deepcopy(head_tpl), header))
    for i, r in enumerate(rows):
        new_rows.append((copy.deepcopy(body_tpl[i % len(body_tpl)]), r))
    for tr, vals in new_rows:
        for ci, tc in enumerate(tr.findall(qn('a:tc'))):
            txt = str(vals[ci]) if ci < len(vals) else ''
            tb = tc.find(qn('a:txBody'))
            ps = tb.findall(qn('a:p'))
            for p in ps[1:]:
                tb.remove(p)
            p = ps[0]
            runs = p.findall(qn('a:r'))
            rpr = copy.deepcopy(runs[0].find(qn('a:rPr'))) if runs and runs[0].find(qn('a:rPr')) is not None else None
            for ch in list(p):
                if ch.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
                    p.remove(ch)
            r = p.makeelement(qn('a:r'), {})
            if rpr is not None:
                r.append(rpr)
            t = r.makeelement(qn('a:t'), {})
            t.text = txt
            r.append(t)
            end = p.find(qn('a:endParaRPr'))
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        tbl.append(tr)
    # высота рамки = сумма строк
    xfrm = frame_el.find(qn('p:xfrm'))
    if xfrm is not None:
        h = sum(int(tr.get('h')) for tr in tbl.findall(qn('a:tr')))
        xfrm.find(qn('a:ext')).set('cy', str(h))


def build(spec, template, out):
    prs = Presentation(str(template))
    originals = list(prs.slides)
    filled = []  # (slide_no, layout_name, idx, text) — для проверки переполнения
    for n, s in enumerate(spec['slides'], 1):
        name = s['layout']
        cfg = LAYOUTS[name]
        layout = _layout(prs, name)
        slide = prs.slides.add_slide(layout)
        _add_slide_number(slide, layout)
        phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        used = set()
        for key, f in cfg['fields'].items():
            v = s.get(key, f.get('default'))
            if f['type'] in ('text', 'list') and v not in (None, '', []):
                ph = phs.get(f['idx'])
                if ph is not None:
                    _set_text(ph, v, as_list=f['type'] == 'list')
                    used.add(f['idx'])
                    filled.append((n, name, f['idx'], ph.text_frame.text))
            elif f['type'] == 'image':
                ph = phs.get(f['idx'])
                if ph is not None:
                    used.add(f['idx'])  # пустой плейсхолдер рисунка оставляем для ручной вставки
                    if v and Path(str(v)).exists():
                        ph.insert_picture(str(v))
            elif f['type'] == 'table' and v:
                tpl = _table_from_example(originals, name)
                if tpl is None:
                    print(f'ВНИМАНИЕ: в шаблоне нет таблицы-образца для макета «{name}», таблица пропущена')
                else:
                    el = copy.deepcopy(tpl)
                    _fill_table(el, v.get('header') or [], v.get('rows') or [])
                    slide.shapes._spTree.append(el)
        if 'items' in cfg:
            it = cfg['items']
            for i, item in enumerate((s.get(it['key']) or [])[:it['max']]):
                for fk, f in it['fields'].items():
                    ph = phs.get(f['idx'][i])
                    if ph is not None and item.get(fk):
                        _set_text(ph, item[fk])
                        used.add(f['idx'][i])
                        filled.append((n, name, f['idx'][i], ph.text_frame.text))
        # незаполненные плейсхолдеры (например, 4-я плитка в «Цифрах») делаем пустыми, прозрачными и скрытыми.
        # Просто удалить нельзя: часть редакторов (LibreOffice) тогда дорисовывает плейсхолдер из макета.
        for idx, ph in phs.items():
            if idx not in used and 'SLIDE_NUMBER' not in str(ph.placeholder_format.type):
                _hide_placeholder(ph)
        if s.get('notes'):
            slide.notes_slide.notes_text_frame.text = str(s['notes'])
    # удаляем слайды-примеры шаблона
    lst = prs.slides._sldIdLst
    for sl in originals:
        for sid in list(lst):
            if prs.part.related_part(sid.rId) is sl.part:
                rid = sid.rId
                lst.remove(sid)
                prs.part.drop_rel(rid)
    prs.save(str(out))
    return filled


# ------------------------------------------------------------------ рендер и проверка переполнения

def render(pptx, outdir):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice or not shutil.which('pdftoppm'):
        print('Рендер пропущен: нет LibreOffice (soffice) или pdftoppm')
        return None, []
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        shutil.copy(pptx, Path(td) / 'in.pptx')
        subprocess.run([soffice, f'-env:UserInstallation=file://{td}/lo', '--headless', '--convert-to', 'pdf',
                        '--outdir', td, f'{td}/in.pptx'], capture_output=True, timeout=600)
        if not (Path(td) / 'in.pdf').exists():
            print('Рендер не удался')
            return None, []
        pdf = outdir / 'slides.pdf'
        shutil.copy(Path(td) / 'in.pdf', pdf)
    for f in outdir.glob('slide-*.png'):
        f.unlink()
    subprocess.run(['pdftoppm', '-r', '60', '-png', str(pdf), str(outdir / 'slide')], timeout=600)
    pngs = sorted(outdir.glob('slide-*.png'), key=lambda p: int(re.findall(r'(\d+)\.png$', p.name)[0]))
    try:
        from PIL import Image, ImageDraw
        ims = [Image.open(p).convert('RGB') for p in pngs]
        w = 420
        ims = [im.resize((w, int(im.height * w / im.width))) for im in ims]
        cols, h = 3, ims[0].height
        sheet = Image.new('RGB', (cols * (w + 10) + 10, ((len(ims) + cols - 1) // cols) * (h + 26) + 10), 'white')
        d = ImageDraw.Draw(sheet)
        for i, im in enumerate(ims):
            x, y = 10 + (i % cols) * (w + 10), 10 + (i // cols) * (h + 26)
            d.text((x, y), f'#{i + 1}', fill='black')
            sheet.paste(im, (x, y + 14))
        sheet.save(outdir / 'contact.jpg', quality=85)
    except Exception as e:  # PIL необязателен
        print(f'Контакт-лист не собран: {e}')
    return pdf, pngs


def overflow_check(pptx, pdf, filled):
    """Слова текста, отрисованные вне рамки своего поля (переполнение)."""
    if not pdf or not shutil.which('pdftotext'):
        return []
    from collections import Counter
    prs = Presentation(str(pptx))
    W = prs.slide_width
    issues = []
    for n, lay_name, idx, text in filled:
        out = subprocess.run(['pdftotext', '-f', str(n), '-l', str(n), '-bbox', str(pdf), '-'],
                             capture_output=True, text=True).stdout
        m = re.search(r'<page width="([\d.]+)" height="([\d.]+)"', out)
        if not m:
            continue
        pw, phh = float(m.group(1)), float(m.group(2))
        k = pw / W
        lay = next(l for l in prs.slide_layouts if l.name == lay_name)
        lph = next((p for p in lay.placeholders if p.placeholder_format.idx == idx), None)
        if lph is None:
            continue
        x0, y0 = lph.left * k, lph.top * k
        x1, y1 = x0 + lph.width * k, y0 + lph.height * k
        tol = 0.02 * phh
        pos = {}
        for w in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>', out):
            pos.setdefault(w.group(5).lower().strip('.,:;!?«»"()—'), []).append(
                ((float(w.group(1)) + float(w.group(3))) / 2, (float(w.group(2)) + float(w.group(4))) / 2))
        need = Counter(t.lower().strip('.,:;!?«»"()—') for t in text.split() if len(t) >= 4)
        miss = tot = 0
        for wd, cnt in need.items():
            if wd not in pos:
                continue
            tot += cnt
            inside = sum(1 for cx, cy in pos[wd] if x0 - tol <= cx <= x1 + tol and y0 - tol <= cy <= y1 + tol)
            miss += max(0, cnt - inside)
        if tot and miss / tot > 0.1:
            issues.append(f'Слайд {n} («{lay_name}»): текст поля «{lph.name}» вылезает за рамку '
                          f'({miss} из {tot} слов снаружи) — сократи')
    return issues


# ------------------------------------------------------------------ CLI

def print_catalog():
    for name, cfg in LAYOUTS.items():
        print(f'\n## {name} — {cfg["use"]}')
        for k, f in cfg['fields'].items():
            lim = f' (до {f["max"]} зн.)' if 'max' in f and f['type'] == 'text' else ''
            if f['type'] == 'list':
                lim = f' (до {f["max_items"]} пунктов по {f["max"]} зн.)'
            print(f'  "{k}": {f["label"]}{lim}{" *" if f.get("required") else ""}')
        if 'items' in cfg:
            it = cfg['items']
            fl = ', '.join(f'"{k}" до {f["max"]} зн.' for k, f in it['fields'].items())
            print(f'  "{it["key"]}": {it["label"]}, от {it["min"]} до {it["max"]} шт., каждый {{{fl}}}')
    print('\nПравила:\n- ' + '\n- '.join(CATALOG['rules']))


def main():
    ap = argparse.ArgumentParser(description='Сборка презентации VITOMEK из корпоративного шаблона')
    ap.add_argument('content', nargs='?', help='content.json')
    ap.add_argument('-o', '--out', help='итоговый .pptx')
    ap.add_argument('--template', default=str(DEFAULT_TEMPLATE))
    ap.add_argument('--render', action='store_true', help='картинки слайдов + проверка переполнения')
    ap.add_argument('--check', action='store_true', help='только проверить content.json')
    ap.add_argument('--catalog', action='store_true', help='показать макеты и лимиты')
    a = ap.parse_args()
    if a.catalog:
        print_catalog()
        return
    if not a.content:
        ap.error('укажи content.json')
    spec = json.loads(Path(a.content).read_text(encoding='utf-8'))
    errors, warns = validate(spec)
    for e in errors:
        print('ОШИБКА: ' + e)
    for w in warns:
        print('ВНИМАНИЕ: ' + w)
    if errors:
        sys.exit('Исправь ошибки в content.json и запусти снова.')
    if a.check:
        print('content.json в порядке' if not warns else 'Собрать можно, но лучше поправить предупреждения.')
        return
    out = Path(a.out or Path(a.content).with_suffix('.pptx'))
    out.parent.mkdir(parents=True, exist_ok=True)
    filled = build(spec, a.template, out)
    print(f'ГОТОВО: {out} ({len(spec["slides"])} слайдов)')
    if a.render:
        rdir = out.parent / (out.stem + '_preview')
        pdf, pngs = render(out, rdir)
        if pngs:
            print(f'Картинки слайдов: {rdir}/slide-NN.png, обзор: {rdir}/contact.jpg')
            issues = overflow_check(out, pdf, filled)
            for i in issues:
                print('ПЕРЕПОЛНЕНИЕ: ' + i)
            if not issues:
                print('Переполнений не найдено')


if __name__ == '__main__':
    main()
